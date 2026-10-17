"""Build, or check, the macro-free workbook template.

The template src/workbook/FRTB_Template.xlsx is the workbook that
M_Setup.BuildWorkbook produces, saved without a VBA project, without sheet
buttons and without personal document properties. Builds start from it and
import the modules of src/modules (and tests/modules); M_Setup.FinishWorkbook
then adds the buttons and window settings a template cannot hold.

The build runs headless in LibreOffice Calc. tools/Build-Workbook.ps1
-Template produces the same content through Excel.

Usage:
    python3 tools/build_template.py          # rewrite the template
    python3 tools/build_template.py --check  # exit 1 if it differs from a fresh build
"""

from __future__ import annotations

import argparse
import re
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "tools"))

import lo_runner  # noqa: E402

TEMPLATE = lo_runner.TEMPLATE


def _strip(doc) -> None:
    """Remove sheet controls and personal document properties."""
    sheets = doc.Sheets
    for i in range(sheets.Count):
        page = sheets.getByIndex(i).DrawPage
        while page.Count > 0:
            page.remove(page.getByIndex(0))
    props = doc.DocumentProperties
    props.Author = ""
    props.ModifiedBy = ""
    props.Title = ""
    props.Subject = ""
    props.Description = ""


def build(output: Path) -> None:
    errors = lo_runner.run(output, ["M_Setup.BuildWorkbook"], post_process=_strip)
    if errors:
        raise RuntimeError("; ".join(errors))


def _color(c) -> str | None:
    if c is None:
        return None
    if c.type == "rgb":
        return str(c.rgb)
    return f"{c.type}:{c.value}:{round(c.tint or 0, 4)}"


def _cell_style(c) -> tuple:
    f, fill, a, b = c.font, c.fill, c.alignment, c.border
    return (
        c.number_format,
        (f.name, f.sz, bool(f.b), bool(f.i), _color(f.color)),
        (fill.fill_type, _color(fill.fgColor) if fill.fill_type else None),
        (a.horizontal, a.vertical, bool(a.wrap_text)),
        tuple((side.style, _color(side.color)) for side in (b.left, b.right, b.top, b.bottom)),
    )


def _attributes(element: str) -> dict:
    return dict(re.findall(r'([\w:]+)="([^"]*)"', element))


def sheet_parts(path: Path) -> dict:
    """Map each sheet name to its XML part inside the package."""
    with zipfile.ZipFile(path) as z:
        book = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = {}
    for element in re.findall(r"<Relationship [^>]*>", rels):
        a = _attributes(element)
        targets[a["Id"]] = "xl/" + a["Target"].lstrip("/").removeprefix("xl/")
    return {a["name"]: targets[a["r:id"]] for a in map(_attributes, re.findall(r"<sheet [^>]*>", book))}


def _explicit_dimensions(path: Path) -> dict:
    """Row heights and column widths set explicitly, per sheet name.

    Automatic heights and widths depend on the fonts installed where the
    workbook was saved, so only entries flagged customHeight/customWidth
    are compared. openpyxl does not keep those flags, so read the XML.
    """
    out = {}
    with zipfile.ZipFile(path) as z:
        for name, part in sheet_parts(path).items():
            xml = z.read(part).decode("utf-8")
            rows, cols = {}, {}
            for element in re.findall(r"<row [^>]*>", xml):
                a = _attributes(element)
                if a.get("customHeight") in ("true", "1") and "ht" in a:
                    rows[int(a["r"])] = float(a["ht"])
            for element in re.findall(r"<col [^>]*>", xml):
                a = _attributes(element)
                if a.get("customWidth") in ("true", "1") and "width" in a:
                    cols[(int(a["min"]), int(a["max"]))] = round(float(a["width"]), 2)
            out[name] = (rows, cols)
    return out


def signature(path: Path) -> dict:
    """Normalized content and layout of a workbook: values, cell styles,
    explicit column widths and row heights, merged ranges, tab colors,
    frozen panes, gridlines and autofilters of every sheet."""
    wb = openpyxl.load_workbook(path)
    dims = _explicit_dimensions(path)
    sheets = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None or c.has_style:
                    cells[c.coordinate] = (c.value, _cell_style(c))
        sheets[ws.title] = {
            "cells": cells,
            "columns": dims.get(ws.title, ({}, {}))[1],
            "rows": dims.get(ws.title, ({}, {}))[0],
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "tab": _color(ws.sheet_properties.tabColor),
            "freeze": ws.freeze_panes,
            "gridlines": ws.sheet_view.showGridLines,
            "autofilter": ws.auto_filter.ref,
        }
    return {"order": wb.sheetnames, "sheets": sheets}


def signature_differences(a: dict, b: dict, limit: int = 5) -> list[str]:
    """Describe the first differences between two signatures."""
    out = []
    if a["order"] != b["order"]:
        out.append(f"sheet order differs: {a['order']} vs {b['order']}")
    for name in sorted(a["sheets"].keys() & b["sheets"].keys()):
        sa, sb = a["sheets"][name], b["sheets"][name]
        for key in ("columns", "rows", "merged", "tab", "freeze", "gridlines", "autofilter"):
            if sa[key] != sb[key]:
                out.append(f"{name}: {key} differs")
        for coord in sorted(sa["cells"].keys() ^ sb["cells"].keys()):
            out.append(f"{name}!{coord}: present in only one workbook")
        for coord in sorted(sa["cells"].keys() & sb["cells"].keys()):
            if sa["cells"][coord] != sb["cells"][coord]:
                out.append(f"{name}!{coord}: {sa['cells'][coord]} vs {sb['cells'][coord]}")
            if len(out) >= limit:
                return out
    return out[:limit]


def package_problems(path: Path) -> list[str]:
    """Return what makes the file unsuitable as a macro-free template."""
    problems = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        if any("vbaProject" in n for n in names):
            problems.append("contains a VBA project")
        if any(n.startswith("xl/ctrlProps/") for n in names):
            problems.append("contains form controls")
        core = z.read("docProps/core.xml").decode("utf-8") if "docProps/core.xml" in names else ""
        for tag in ("dc:creator", "cp:lastModifiedBy"):
            start = core.find(f"<{tag}>")
            if start >= 0 and core[start + len(tag) + 2 : core.find(f"</{tag}>")].strip():
                problems.append(f"document property {tag} is not empty")
    return problems


def check() -> list[str]:
    if not TEMPLATE.exists():
        return [f"{TEMPLATE} does not exist"]
    problems = package_problems(TEMPLATE)
    with tempfile.TemporaryDirectory() as tmp:
        fresh = Path(tmp) / "fresh.xlsx"
        build(fresh)
        diffs = signature_differences(signature(TEMPLATE), signature(fresh))
    problems.extend(f"differs from a fresh build: {d}" for d in diffs)
    return problems


def main(argv) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    if args.check:
        problems = check()
        for p in problems:
            print("TEMPLATE:", p, file=sys.stderr)
        if not problems:
            print("Template is current and macro-free")
        return 1 if problems else 0
    TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
    build(TEMPLATE)
    problems = package_problems(TEMPLATE)
    for p in problems:
        print("TEMPLATE:", p, file=sys.stderr)
    print(f"Wrote {TEMPLATE.relative_to(REPO_ROOT)}")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
