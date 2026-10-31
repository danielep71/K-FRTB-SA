"""Hand-calculated checks of the reference implementation.

Each expected value is written out from the MAR21-MAR23 formulas with the
numbers of a small, explicit example, so that the reference (and through
reconciliation the VBA engine) is tied to arithmetic that can be followed
on paper.
"""

import math

import pytest

import frtb_reference as ref
from regulatory_defaults import build_params


@pytest.fixture
def p():
    return build_params()


def run_delta(p, rows):
    capital, alt, buckets = ref.sbm_component(p, "DELTA", ref.delta_risk_factors(p, rows))
    return capital, alt, buckets


def approx(x):
    return pytest.approx(x, rel=1e-10, abs=1e-8)


# --------------------------------------------------------------------------
# Correlation scenarios (MAR21.6)
# --------------------------------------------------------------------------
@pytest.mark.parametrize(
    "rho, low, high",
    [(0.5, 0.375, 0.625), (0.9, 0.8, 1.0), (0.2, 0.15, 0.25), (1.0, 1.0, 1.0), (0.0, 0.0, 0.0)],
)
def test_scenario_transform(p, rho, low, high):
    assert float(ref.scenario(rho, "LOW", p)) == approx(low)
    assert float(ref.scenario(rho, "MEDIUM", p)) == approx(rho)
    assert float(ref.scenario(rho, "HIGH", p)) == approx(high)


# --------------------------------------------------------------------------
# GIRR delta
# --------------------------------------------------------------------------
def test_girr_same_curve_two_tenors(p):
    rows = [["1", "GIRR", "EUR", "ESTR", "YIELD", 1, 1_000_000], ["2", "GIRR", "EUR", "ESTR", "YIELD", 5, -500_000]]
    capital, _, _ = run_delta(p, rows)
    ws1, ws2 = 0.016 * 1_000_000, 0.011 * -500_000
    rho = max(math.exp(-0.03 * 4 / 1), 0.40)
    k = lambda r: math.sqrt(ws1**2 + ws2**2 + 2 * r * ws1 * ws2)  # noqa: E731
    assert capital["GIRR"]["MEDIUM"] == approx(k(rho))
    assert capital["GIRR"]["HIGH"] == approx(10_500.0)  # rho_high = 1
    assert capital["GIRR"]["LOW"] == approx(k(max(2 * rho - 1, 0.75 * rho)))


def test_girr_tenor_correlation_floor(p):
    rows = [["1", "GIRR", "EUR", "ESTR", "YIELD", 0.25, 1_000_000], ["2", "GIRR", "EUR", "ESTR", "YIELD", 30, 1_000_000]]
    capital, _, _ = run_delta(p, rows)
    ws1, ws2 = 17_000.0, 11_000.0
    # exp(-0.03 * 29.75 / 0.25) is far below the 40% floor
    assert capital["GIRR"]["MEDIUM"] == approx(math.sqrt(ws1**2 + ws2**2 + 2 * 0.40 * ws1 * ws2))


def test_girr_different_curves_inflation_and_xccy(p):
    rows = [
        ["1", "GIRR", "EUR", "ESTR", "YIELD", 5, 1_000_000],
        ["2", "GIRR", "EUR", "EURIBOR6M", "YIELD", 5, 1_000_000],
        ["3", "GIRR", "EUR", "HICP", "INFLATION", None, 1_000_000],
        ["4", "GIRR", "EUR", "EUR-USD", "XCCY", None, 1_000_000],
    ]
    capital, _, _ = run_delta(p, rows)
    y, i, x = 11_000.0, 16_000.0, 16_000.0
    k2 = 2 * y**2 + i**2 + x**2 + 2 * 0.999 * y * y + 2 * 2 * 0.40 * y * i
    assert capital["GIRR"]["MEDIUM"] == approx(math.sqrt(k2))


def test_girr_two_currencies(p):
    rows = [["1", "GIRR", "EUR", "ESTR", "YIELD", 10, 1_000_000], ["2", "GIRR", "USD", "SOFR", "YIELD", 10, -1_000_000]]
    capital, alt, _ = run_delta(p, rows)
    k = 11_000.0
    assert capital["GIRR"]["MEDIUM"] == approx(math.sqrt(2 * k**2 - 2 * 0.50 * k * k))
    assert capital["GIRR"]["LOW"] == approx(math.sqrt(2 * k**2 - 2 * 0.375 * k * k))
    assert capital["GIRR"]["HIGH"] == approx(math.sqrt(2 * k**2 - 2 * 0.625 * k * k))
    assert not any(alt["GIRR"].values())


def test_girr_netting_of_identical_risk_factors(p):
    rows = [["1", "GIRR", "USD", "SOFR", "YIELD", 10, 2_000_000], ["2", "GIRR", "USD", "SOFR", "YIELD", 10, 500_000]]
    capital, _, buckets = run_delta(p, rows)
    assert capital["GIRR"]["MEDIUM"] == approx(0.011 * 2_500_000)
    assert buckets[("GIRR", "MEDIUM", "USD")][2] == 1


def test_girr_specified_currency_discretion(p):
    p.girr_specified_on = True
    rows = [["1", "GIRR", "EUR", "ESTR", "YIELD", 10, 1_000_000], ["2", "GIRR", "PLN", "WIBOR", "YIELD", 10, 1_000_000]]
    rfs = ref.delta_risk_factors(p, rows)
    assert rfs[0].rw == approx(0.011 / math.sqrt(2))
    assert rfs[1].rw == approx(0.011)
    # MAR21.44 divides "the above risk weights": inflation and cross-currency basis too (#29)
    rows = [["1", "GIRR", "EUR", "HICP", "INFLATION", None, 1], ["2", "GIRR", "USD", "EUR", "XCCY", None, 1],
            ["3", "GIRR", "PLN", "CPI", "INFLATION", None, 1]]
    rw = [f.rw for f in ref.delta_risk_factors(p, rows)]
    assert rw == [approx(0.016 / math.sqrt(2)), approx(0.016 / math.sqrt(2)), approx(0.016)]
    p.girr_specified_on = False
    infl = ref.delta_risk_factors(p, [["1", "GIRR", "EUR", "HICP", "INFLATION", None, 1]])
    assert infl[0].rw == approx(0.016)


# --------------------------------------------------------------------------
# CSR delta: alternative Sb and the other-sector bucket
# --------------------------------------------------------------------------
CSR_ROWS = [
    ["1", "CSR_NS", 1, "ITALY", "BOND", 5, 3_000_000],
    ["2", "CSR_NS", 1, "SPAIN", "BOND", 5, 3_000_000],
    ["3", "CSR_NS", 2, "REGION_X", "BOND", 3, -2_000_000],
    ["4", "CSR_NS", 2, "REGION_Y", "BOND", 5, -1_800_000],
]


def test_csr_alternative_sb(p):
    capital, alt, buckets = run_delta(p, CSR_ROWS)
    k1 = math.sqrt(2 * 15_000**2 + 2 * 0.35 * 15_000**2)
    k2 = math.sqrt(20_000**2 + 18_000**2 + 2 * 0.35 * 0.65 * 20_000 * 18_000)
    s1, s2 = 30_000.0, -38_000.0
    assert k1**2 + k2**2 + 2 * 0.75 * s1 * s2 < 0  # radicand negative: MAR21.4(5)(b)
    s1a, s2a = min(s1, k1), max(s2, -k2)
    assert alt["CSR_NS"]["MEDIUM"]
    assert capital["CSR_NS"]["MEDIUM"] == approx(math.sqrt(k1**2 + k2**2 + 2 * 0.75 * s1a * s2a))
    assert buckets[("CSR_NS", "MEDIUM", "1")][1] == approx(s1a)
    assert buckets[("CSR_NS", "MEDIUM", "1")][3] == "ALT_SB"


def test_csr_other_sector_bucket_added_without_diversification(p):
    rows = CSR_ROWS[:2] + [
        ["5", "CSR_NS", 16, "MISC_CO", "BOND", 5, 200_000],
        ["6", "CSR_NS", 16, "MISC_CO2", "BOND", 5, -150_000],
    ]
    capital, _, buckets = run_delta(p, rows)
    k1 = math.sqrt(2 * 15_000**2 + 2 * 0.35 * 15_000**2)
    assert buckets[("CSR_NS", "MEDIUM", "16")][0] == approx(24_000 + 18_000)
    assert capital["CSR_NS"]["MEDIUM"] == approx(k1 + 42_000)
    assert capital["CSR_NS"]["HIGH"] == approx(15_000 * math.sqrt(2 + 2 * 0.4375) + 42_000)


def test_csr_bond_cds_basis_and_index_bucket(p):
    rows = [["1", "CSR_NS", 3, "BANK_A", "BOND", 3, 1_500_000], ["2", "CSR_NS", 3, "BANK_A", "CDS", 3, -1_400_000]]
    capital, _, _ = run_delta(p, rows)
    a, b = 75_000.0, -70_000.0
    assert capital["CSR_NS"]["MEDIUM"] == approx(math.sqrt(a * a + b * b + 2 * 0.999 * a * b))
    rows = [["1", "CSR_NS", 17, "ITRAXX_MAIN", "CDS", 5, 1_000_000], ["2", "CSR_NS", 17, "CDX_IG", "CDS", 5, 1_000_000]]
    capital, _, _ = run_delta(p, rows)
    assert capital["CSR_NS"]["MEDIUM"] == approx(15_000 * math.sqrt(2 + 2 * 0.80))


# --------------------------------------------------------------------------
# Equity, commodity and FX delta
# --------------------------------------------------------------------------
def test_equity_spot_repo_and_names(p):
    rows = [
        ["1", "EQ", 8, "BANK_A", "SPOT", None, 2_000_000],
        ["2", "EQ", 8, "BANK_B", "SPOT", None, -1_200_000],
        ["3", "EQ", 8, "BANK_A", "REPO", None, 50_000],
    ]
    capital, _, _ = run_delta(p, rows)
    a, b, r = 1_000_000.0, -600_000.0, 250.0
    k2 = a * a + b * b + r * r + 2 * (0.25 * a * b + 0.999 * a * r + 0.25 * 0.999 * b * r)
    assert capital["EQ"]["MEDIUM"] == approx(math.sqrt(k2))


def test_equity_cross_bucket_index(p):
    rows = [["1", "EQ", 8, "BANK_A", "SPOT", None, 2_000_000], ["2", "EQ", 12, "SX5E", "SPOT", None, -1_500_000]]
    capital, _, _ = run_delta(p, rows)
    a, b = 1_000_000.0, -225_000.0
    assert capital["EQ"]["MEDIUM"] == approx(math.sqrt(a * a + b * b + 2 * 0.45 * a * b))


def test_commodity_tenor_and_location(p):
    rows = [
        ["1", "COM", 2, "BRENT", "ROTTERDAM", 0.5, 800_000],
        ["2", "COM", 2, "BRENT", "ROTTERDAM", 1, -500_000],
        ["3", "COM", 2, "WTI", "CUSHING", 1, 300_000],
    ]
    capital, _, _ = run_delta(p, rows)
    a, b, c = 280_000.0, -175_000.0, 105_000.0
    k2 = a * a + b * b + c * c + 2 * (0.99 * a * b + 0.95 * 0.99 * 0.999 * a * c + 0.95 * 0.999 * b * c)
    assert capital["COM"]["MEDIUM"] == approx(math.sqrt(k2))


def test_fx_buckets_and_specified_pairs(p):
    rows = [["1", "FX", "USD", "USD", None, None, 3_000_000], ["2", "FX", "GBP", "GBP", None, None, -1_500_000]]
    capital, _, _ = run_delta(p, rows)
    a, b = 450_000.0, -225_000.0
    assert capital["FX"]["MEDIUM"] == approx(math.sqrt(a * a + b * b + 2 * 0.60 * a * b))
    p.fx_specified_on = True
    rfs = ref.delta_risk_factors(p, rows + [["3", "FX", "PLN", "PLN", None, None, 1]])
    rw = {f.bucket: f.rw for f in rfs}
    assert rw["USD"] == approx(0.15 / math.sqrt(2))  # USD/EUR is specified
    assert rw["GBP"] == approx(0.15 / math.sqrt(2))  # EUR/GBP is specified
    assert rw["PLN"] == approx(0.15)


def test_fx_reporting_currency_is_not_a_risk_factor(p):
    assert ref.delta_risk_factors(p, [["1", "FX", "EUR", "EUR", None, None, 1_000]]) == []


# --------------------------------------------------------------------------
# Vega
# --------------------------------------------------------------------------
def test_vega_risk_weights(p):
    assert ref.vega_rw(p, "GIRR", "EUR") == approx(1.0)
    assert ref.vega_rw(p, "FX", "USD") == approx(1.0)
    assert ref.vega_rw(p, "EQ", "8") == approx(0.55 * math.sqrt(2))
    assert ref.vega_rw(p, "EQ", "12") == approx(0.55 * math.sqrt(2))
    assert ref.vega_rw(p, "EQ", "13") == approx(0.55 * math.sqrt(2))  # indices count as large cap (FAQ)
    assert ref.vega_rw(p, "EQ", "9") == approx(1.0)
    assert ref.vega_rw(p, "COM", "2") == approx(1.0)


def test_vega_option_maturity_correlation(p):
    rows = [["1", "EQ", 8, "BANK_A", None, 0.5, None, 300_000], ["2", "EQ", 8, "BANK_A", None, 1, None, -100_000]]
    capital, _, _ = ref.sbm_component(p, "VEGA", ref.vega_risk_factors(p, rows))
    rw = 0.55 * math.sqrt(2)
    a, b = rw * 300_000, rw * -100_000
    rho = math.exp(-0.01 * 0.5 / 0.5)
    assert capital["EQ"]["MEDIUM"] == approx(math.sqrt(a * a + b * b + 2 * rho * a * b))


def test_vega_girr_underlying_maturity(p):
    rows = [["1", "GIRR", "EUR", None, None, 1, 5, 400_000], ["2", "GIRR", "EUR", None, None, 5, 10, -250_000]]
    capital, _, _ = ref.sbm_component(p, "VEGA", ref.vega_risk_factors(p, rows))
    rho = math.exp(-0.01 * 4 / 1) * math.exp(-0.01 * 5 / 5)
    a, b = 400_000.0, -250_000.0
    assert capital["GIRR"]["MEDIUM"] == approx(math.sqrt(a * a + b * b + 2 * rho * a * b))


# --------------------------------------------------------------------------
# Curvature
# --------------------------------------------------------------------------
def run_curv(p, rows):
    return ref.sbm_component(p, "CURVATURE", ref.curvature_risk_factors(p, rows))


def test_curvature_bucket_selection_and_psi(p):
    rows = [["1", "EQ", 8, "BANK_A", None, -80_000, 40_000], ["2", "EQ", 8, "BANK_B", None, 20_000, -60_000]]
    capital, _, buckets = run_curv(p, rows)
    rho = 0.25**2
    k_up = math.sqrt(0 + 20_000**2 + 2 * rho * -80_000 * 20_000)
    k_dn = math.sqrt(40_000**2 + 0 + 2 * rho * 40_000 * -60_000)
    assert k_dn > k_up
    assert capital["EQ"]["MEDIUM"] == approx(k_dn)
    assert buckets[("EQ", "MEDIUM", "8")][1] == approx(-20_000)
    assert buckets[("EQ", "MEDIUM", "8")][3] == "DOWN"


def test_curvature_both_negative_has_no_cross_term(p):
    rows = [["1", "EQ", 8, "A", None, -10_000, 0], ["2", "EQ", 8, "B", None, -20_000, 0]]
    _, _, buckets = run_curv(p, rows)
    # K+ = 0 (psi removes the cross term); K- = 0; tie broken by the larger sum (0 > -30000)
    kb, sb, _, note = buckets[("EQ", "MEDIUM", "8")]
    assert kb == approx(0.0) and sb == approx(0.0) and note == "DOWN"


def test_curvature_across_buckets_floor_at_zero(p):
    rows = [["1", "FX", "USD", "USD", None, -40_000, -35_000], ["2", "FX", "GBP", "GBP", None, 15_000, 12_000]]
    capital, _, buckets = run_curv(p, rows)
    assert buckets[("FX", "MEDIUM", "USD")][1] == approx(-35_000)  # tie at K = 0 -> larger sum
    assert 15_000**2 + 2 * 0.36 * -35_000 * 15_000 < 0
    assert capital["FX"]["MEDIUM"] == approx(0.0)


def test_curvature_fx_divisor(p):
    # MAR21.98 (#29, item 3): the discretion divides options that do not reference
    # the reporting currency; the approved alternative (_ALL) divides every FX row
    rows = [["1", "FX", "USD", "USD", None, 30_000, 0, None], ["2", "FX", "USD", "USD", None, 30_000, 0, True]]
    assert run_curv(p, rows)[0]["FX"]["MEDIUM"] == approx(60_000)
    p.fx_curv_divisor_on = True
    assert run_curv(p, rows)[0]["FX"]["MEDIUM"] == approx(20_000 + 30_000)
    p.fx_curv_divisor_all_on = True
    assert run_curv(p, rows)[0]["FX"]["MEDIUM"] == approx(40_000)
    p.fx_curv_divisor_on = False
    assert run_curv(p, rows)[0]["FX"]["MEDIUM"] == approx(40_000)


@pytest.mark.parametrize("flag", ["maybe", "YES", "y", "No", "N", "1", "-1", "0", 1, -1, 0, 2, 0.5])
def test_curvature_reporting_ccy_flag_rejects_other_values(p, flag):
    # Only TRUE, FALSE or blank are accepted; an alias skips the row (#144)
    rows = [["1", "FX", "USD", "USD", None, 30_000, 0, flag], ["2", "FX", "GBP", "GBP", None, 10_000, 0, None]]
    capital, _, _ = run_curv(p, rows)
    assert capital["FX"]["MEDIUM"] == approx(10_000)


@pytest.mark.parametrize("flag, divided", [(None, True), ("", True), ("  ", True), (False, True), ("false", True),
                                           (" FALSE ", True), (True, False), ("TRUE", False), ("true", False)])
def test_curvature_reporting_ccy_flag_accepts_true_false_or_blank(p, flag, divided):
    p.fx_curv_divisor_on = True
    rows = [["1", "FX", "USD", "USD", None, 30_000, 0, flag]]
    capital, _, _ = run_curv(p, rows)
    assert capital["FX"]["MEDIUM"] == approx(20_000 if divided else 30_000)


def test_curvature_other_bucket(p):
    rows = [["1", "EQ", 11, "A", None, 10_000, -5_000], ["2", "EQ", 11, "B", None, 4_000, 7_000]]
    capital, _, _ = run_curv(p, rows)
    assert capital["EQ"]["MEDIUM"] == approx(max(14_000, 7_000))


# --------------------------------------------------------------------------
# DRC
# --------------------------------------------------------------------------
def drc_row(obl, bucket, quality, seniority, notional, mv, maturity=1, lgd=None):
    return ["x", obl, bucket, quality, seniority, notional, mv, maturity, lgd]


def test_drc_jtd_long_short_and_maturity(p):
    total, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "SENIOR", 100, 95)])
    assert obl["A"][3] == approx(0.75 * 100 + (95 - 100))
    _, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "SENIOR", -100, -95)])
    assert obl["A"][4] == approx(70)
    _, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "SENIOR", 100, 95, maturity=0.1)])
    assert obl["A"][3] == approx(70 * 0.25)
    _, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "SENIOR", 100, 95, maturity=0.5)])
    assert obl["A"][3] == approx(70 * 0.5)
    assert total == approx(0.06 * 70)


def test_drc_seniority_offsetting(p):
    # A senior short does not hedge an equity long ...
    _, obl, _ = ref.drc(p, [drc_row("A", "CORP", "A", "EQUITY", 100, 100), drc_row("A", "CORP", "A", "SENIOR", -100, -100)])
    assert obl["A"][3] == approx(100) and obl["A"][4] == approx(75)
    # ... but an equity short hedges a senior long
    _, obl, _ = ref.drc(p, [drc_row("A", "CORP", "A", "SENIOR", 100, 100), drc_row("A", "CORP", "A", "EQUITY", -60, -60)])
    assert obl["A"][3] == approx(15) and obl["A"][4] == approx(0)


def test_drc_hedge_benefit_ratio(p):
    rows = [drc_row("A", "CORP", "BBB", "EQUITY", 100, 100), drc_row("B", "CORP", "AA", "EQUITY", -50, -50)]
    total, _, buckets = ref.drc(p, rows)
    hbr = 100 / 150
    assert buckets["CORP"][2] == approx(hbr)
    assert total == approx(max(0.06 * 100 - hbr * 0.02 * 50, 0))


def test_drc_buckets_do_not_offset(p):
    rows = [drc_row("A", "CORP", "BBB", "EQUITY", 100, 100), drc_row("B", "SOV", "BBB", "SENIOR", -100, -100)]
    total, _, _ = ref.drc(p, rows)
    assert total == approx(6.0)


def test_drc_obligor_split_across_buckets_is_rejected(p):
    # MAR22.22 (#29, item 4a): every position of a split obligor is excluded
    rows = [drc_row("A", "CORP", "BBB", "SENIOR", 100, 100), drc_row("A", "SOV", "BBB", "SENIOR", -50, -50),
            drc_row("B", "CORP", "BBB", "SENIOR", 100, 100)]
    total, obl, buckets = ref.drc(p, rows)
    assert list(obl) == ["B"]
    assert total == approx(0.06 * 75)
    assert buckets["SOV"][:2] == (0, 0)


def test_drc_obligor_takes_the_worst_credit_quality(p):
    # MAR22.24 (#29, item 4b): the highest risk weight, in either order
    for first, second in (("A", "BBB"), ("BBB", "A")):
        rows = [drc_row("A", "CORP", first, "SENIOR", 100, 100), drc_row("A", "CORP", second, "SENIOR", 100, 100)]
        total, obl, _ = ref.drc(p, rows)
        assert obl["A"][1:3] == ("BBB", approx(0.06))
        assert total == approx(0.06 * 150)
    # Equal risk weights keep the first quality
    rows = [drc_row("A", "CORP", "UNRATED", "SENIOR", 100, 100), drc_row("A", "CORP", "BB", "SENIOR", 100, 100)]
    assert ref.drc(p, rows)[1]["A"][1] == "UNRATED"


# --------------------------------------------------------------------------
# RRAO
# --------------------------------------------------------------------------
def test_rrao(p):
    rows = [["1", "EXOTIC", 5_000_000], ["2", "OTHER", 20_000_000], ["3", "UNKNOWN", 1_000]]
    assert ref.rrao(p, rows) == approx(50_000 + 20_000)


# --------------------------------------------------------------------------
# Points confirmed against the BIS text (MAR21.57 Table 5, MAR21.88, MAR21.94, MAR21.100)
# --------------------------------------------------------------------------
@pytest.mark.parametrize(
    "b, c, expected",
    [
        ("1", "2", 0.75),  # same rating, sectors 1 and 2
        ("1", "9", 0.50),  # same sector, IG vs HY within 1-15
        ("2", "11", 0.5 * 0.05),  # rating x sector
        ("8", "12", 0.5 * 0.05),  # covered bonds (IG) vs HY bucket
        ("1", "17", 0.45),  # no rating factor outside 1-15
        ("12", "18", 0.45),
        ("17", "18", 0.75),
    ],
)
def test_csr_cross_bucket_gamma(p, b, c, expected):
    assert ref.gamma(p, "CSR_NS", b, c) == approx(expected)
    assert ref.gamma(p, "CSR_NS", c, b) == approx(expected)


def test_csr_sector_table_bucket_16_column(p):
    assert all(p.csr_sector[("16", g)] == 0 for g in ["1", "2", "3", "4", "5", "6", "7", "8", "17", "18"])


def test_fx_first_order_crosses(p):
    p.fx_specified_on = True
    rows = [["1", "FX", c, c, None, None, 1_000] for c in ("AUD", "NZD", "PLN")]
    rw = {f.bucket: f.rw for f in ref.delta_risk_factors(p, rows)}
    assert rw["AUD"] == approx(0.15 / math.sqrt(2))  # EUR/AUD: cross of USD/EUR and USD/AUD
    assert rw["NZD"] == approx(0.15 / math.sqrt(2))
    assert rw["PLN"] == approx(0.15)


def test_csr_vega_uses_name_correlation_only(p):
    rows = [
        ["1", "CSR_NS", 1, "ITALY", "BOND", 1, None, 100_000],
        ["2", "CSR_NS", 1, "ITALY", "CDS", 1, None, -40_000],
        ["3", "CSR_NS", 1, "SPAIN", "BOND", 3, None, 50_000],
    ]
    rfs = ref.vega_risk_factors(p, rows)
    assert len(rfs) == 2  # bond and CDS vega of one issuer are one risk factor
    capital, _, _ = ref.sbm_component(p, "VEGA", rfs)
    rho = 0.35 * math.exp(-0.01 * 2 / 1)  # name x option maturity, no basis or tenor
    a, b = 60_000.0, 50_000.0
    assert capital["CSR_NS"]["MEDIUM"] == approx(math.sqrt(a * a + b * b + 2 * rho * a * b))


def test_csr_and_commodity_curvature_risk_factors(p):
    rows = [
        ["1", "CSR_NS", 1, "ITALY", "BOND", 10_000, -2_000],
        ["2", "CSR_NS", 1, "ITALY", "CDS", -4_000, 1_000],
        ["3", "COM", 2, "BRENT", "ROTTERDAM", 5_000, 0],
        ["4", "COM", 2, "BRENT", "CUSHING", 1_000, 0],
        ["5", "COM", 2, "WTI", "CUSHING", 3_000, 0],
    ]
    rfs = ref.curvature_risk_factors(p, rows)
    assert len([f for f in rfs if f.rc == "CSR_NS"]) == 1  # bond and CDS curves together
    capital, _, _ = ref.sbm_component(p, "CURVATURE", rfs)
    assert capital["CSR_NS"]["MEDIUM"] == approx(6_000)
    rho = 0.95**2  # commodity correlation squared; delivery location is not a curvature dimension
    assert capital["COM"]["MEDIUM"] == approx(math.sqrt(6_000**2 + 3_000**2 + 2 * rho * 6_000 * 3_000))


# --------------------------------------------------------------------------
# Invalid optional numerics and fractional buckets are rejected, not coerced
# --------------------------------------------------------------------------
def test_fractional_bucket_is_rejected(p):
    assert ref.delta_risk_factors(p, [["1", "EQ", 8.6, "A", "SPOT", None, 1_000]]) == []
    assert ref.delta_risk_factors(p, [["1", "EQ", "8.6", "A", "SPOT", None, 1_000]]) == []
    assert len(ref.delta_risk_factors(p, [["1", "EQ", 8.0, "A", "SPOT", None, 1_000]])) == 1


def test_nonnumeric_tenor_is_rejected(p):
    assert ref.delta_risk_factors(p, [["1", "COM", 2, "BRENT", "X", "abc", 1_000]]) == []
    assert ref.vega_risk_factors(p, [["1", "EQ", 8, "A", None, "abc", None, 1_000]]) == []


def test_drc_nonnumeric_maturity_or_lgd_is_rejected(p):
    total, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "SENIOR", 100, 95, maturity="abc")])
    assert total == 0 and obl == {}
    total, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "SENIOR", 100, 95, lgd="x")])
    assert total == 0 and obl == {}
    _, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "SENIOR", 100, 95, maturity=None)])
    assert obl["A"][3] == approx(70)  # blank maturity: one year


def test_drc_cash_equity_maturity_is_three_months_or_one_year(p):
    # MAR22.16 and its FAQ: no cash-equity maturity between three months and one year.
    # Equity LGD is 100%: JTD = 1.0 x 100 + (95 - 100) = 95.
    for maturity, scale in [(0.25, 0.25), (1, 1), (3, 1), (None, 1)]:
        _, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "EQUITY", 100, 95, maturity=maturity)])
        assert obl["A"][3] == approx(95 * scale), maturity
    for maturity in [0.1, 0.5, 0.99]:
        total, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "EQUITY", 100, 95, maturity=maturity)])
        assert total == 0 and obl == {}, maturity
    # The rule applies to cash equity only: a senior bond at 0.5 years is still scaled.
    _, obl, _ = ref.drc(p, [drc_row("A", "CORP", "BBB", "SENIOR", 100, 95, maturity=0.5)])
    assert obl["A"][3] == approx(70 * 0.5)


def test_input_reader_reads_a_boolean_formatted_number_as_the_flag_it_shows():
    # LibreOffice saves TRUE as 1 with a TRUE/FALSE format; a plain 1 is no flag (#144)
    import openpyxl

    ws = openpyxl.Workbook().active
    cells = [(1, '"TRUE";"TRUE";"FALSE"'), (0, '"TRUE";"TRUE";"FALSE"'), (1, "General"), (0, "#,##0.####"), ("TRUE", "@")]
    for i, (v, fmt) in enumerate(cells, start=1):
        ws.cell(1, i).value = v
        ws.cell(1, i).number_format = fmt
    assert [ref._cell_value(ws.cell(1, i)) for i in range(1, 6)] == [True, False, 1, 0, "TRUE"]
