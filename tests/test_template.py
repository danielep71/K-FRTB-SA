"""The committed workbook template and the build that starts from it."""

from __future__ import annotations

import build_template
import lo_runner
import re
import zipfile

import openpyxl
from conftest import require_libreoffice
from test_vba_engine import FIRST_COL, log_rows, reconcile


def test_template_is_current_and_macro_free():
    require_libreoffice()
    assert build_template.check() == []


def test_source_build_print_areas_match_template(tmp_path):
    """VBA source builds set print extents without the OOXML normalizer."""
    require_libreoffice()
    out = tmp_path / "source-print-areas.xlsx"
    assert lo_runner.run(out, ["M_Setup.BuildWorkbook"]) == []
    source = openpyxl.load_workbook(out)
    template = openpyxl.load_workbook(build_template.TEMPLATE)
    for ws in source:
        assert str(ws.print_area) == str(template[ws.title].print_area), ws.title


def test_template_saved_presentation():
    """The XLSX must look complete before any macro can run."""
    wb = openpyxl.load_workbook(build_template.TEMPLATE)
    for ws in wb:
        assert ws.sheet_view.showGridLines is False, ws.title
        expected_tab = "FF1F3864" if ws.title == "README" else "FF548235" if ws.title == "Test_Expected" else "FF2E75B6" if ws.title in ("Results", "Log") or ws.title.startswith("Out_") else "FFBF9000"
        assert ws.sheet_properties.tabColor.rgb == expected_tab, ws.title
        assert ws.freeze_panes == ("A4" if ws.title == "README" else "D7" if ws.title in ("Control", "Test_Expected") or ws.title.startswith("P_") else "D6"), ws.title
        assert ws.page_setup.orientation == "landscape", ws.title
        assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
        assert ws.sheet_properties.pageSetUpPr.fitToPage
        assert "Aptos" in str(ws.oddFooter)
    for name in ("Control", "In_Delta", "In_Vega", "In_Curvature", "In_DRC", "In_RRAO"):
        ws = wb[name]
        row = 7 if name == "Control" else 6
        assert ws.row_dimensions[row].height == 20.1, name
        assert ws.cell(row, 4).border.bottom.style == "thin", name
    assert wb["README"]["D7"].border.bottom.style == "dashed"
    assert wb["README"]["C6"].border.top.style == "thin"
    assert wb["README"]["C6"].border.right.style == "dashed"
    assert wb["README"]["D12"].border.bottom.style == "thin"
    assert wb["In_Delta"]["B2"].value == "Delta sensitivities"
    assert wb["In_Delta"]["D5"].value == "Risk class"
    for name in ("Results", "Out_Summary", "Out_RiskClass", "Out_Buckets", "Out_DRC", "Log"):
        assert wb[name]["C5"].fill.fgColor.rgb == "FF2E75B6", name
        assert wb[name]["C6"].value == "No results yet: press Run FRTB-SA.", name


def test_presentation_normalization_preserves_content_and_is_idempotent(tmp_path):
    import shutil
    import template_presentation

    path = tmp_path / "template.xlsx"
    shutil.copyfile(build_template.TEMPLATE, path)
    before = build_template.signature(path)
    template_presentation.normalize(path)
    after = build_template.signature(path)
    assert build_template.signature_differences(before, after) == []
    template_presentation.normalize(path)
    assert build_template.signature_differences(after, build_template.signature(path)) == []


def _repack(source, target, part, old, new):
    """Copy a package, replacing one occurrence of old by new in one part."""
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part and old:
                text = data.decode("utf-8")
                assert old in text, (part, old)
                data = text.replace(old, new, 1).encode("utf-8")
            zout.writestr(item, data)


def test_signature_detects_layout_changes(tmp_path):
    """A width-, height-, tab-, fill- or format-only change must make the template stale."""
    template = build_template.TEMPLATE
    base = build_template.signature(template)
    control = build_template.sheet_parts(template)["Control"]
    control_xml = zipfile.ZipFile(template).read(control).decode("utf-8")
    tab_color = re.search(r"<tabColor[^>]*/>", control_xml).group(0)

    unchanged = tmp_path / "unchanged.xlsx"
    _repack(template, unchanged, control, "", "")
    assert build_template.signature_differences(base, build_template.signature(unchanged)) == []

    cases = {
        "width": (control, 'width="34.75"', 'width="60"'),
        "height": (control, 'ht="26.1"', 'ht="40"'),
        "tab": (control, tab_color, '<tabColor rgb="FF00FF00"/>'),
        "gridlines": (control, 'showGridLines="0"', 'showGridLines="1"'),
        "freeze": (control, 'topLeftCell="D7"', 'topLeftCell="D8"'),
        "print": (control, 'orientation="landscape"', 'orientation="portrait"'),
        "fill": ("xl/styles.xml", "FF265A99", "FF265A98"),
        "number_format": ("xl/styles.xml", 'formatCode="0.0##%"', 'formatCode="0.00%"'),
    }
    for name, (part, old, new) in cases.items():
        path = tmp_path / f"{name}.xlsx"
        _repack(template, path, part, old, new)
        assert build_template.signature_differences(base, build_template.signature(path)), name


def test_build_from_template_runs_and_reconciles(tmp_path):
    require_libreoffice()
    out = tmp_path / "from-template.xlsx"
    errors = lo_runner.run(
        out,
        ["M_Setup.FinishWorkbook", "M_Main.RunFRTB_Silent", "TEST_Regression.Test_RunRegression_Silent"],
        input_book=build_template.TEMPLATE,
    )
    assert errors == []
    wb = openpyxl.load_workbook(out, data_only=True)
    assert str(wb["Test_Expected"].cell(3, FIRST_COL).value).startswith("Result: PASS")
    assert log_rows(wb, "ERROR") == []
    assert reconcile(wb) == []
