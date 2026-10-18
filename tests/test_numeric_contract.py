"""Synthetic runtime regressions for issues #41 and #45; no locale-dependent expectations."""

import openpyxl
import pytest

import frtb_reference as ref
import lo_runner
from conftest import require_libreoffice
from test_vba_engine import FIRST_COL, log_rows, reconcile, sheet_rows


def test_shared_numeric_validation_and_conversion(tmp_path):
    require_libreoffice()
    assert lo_runner.run(
        tmp_path / "numeric-contract.xlsx",
        ["TEST_Numeric.Test_NumericContracts_Silent"],
        expected_results={"TEST_Numeric.Test_NumericContracts_Silent": 1},
    ) == []


@pytest.mark.parametrize("invalid", ["abc", "", "0,5", "0.5junk", "1e309"])
def test_required_parameter_aborts_and_logs_context(lo_default_path, tmp_path, invalid):
    wb = openpyxl.load_workbook(lo_default_path)
    ws = wb["P_FX"]
    row = next(row for row in ws.iter_rows() if row[2].value == "RW")
    row[3].value = invalid
    source = tmp_path / "bad-parameter.xlsx"
    wb.save(source)
    out = tmp_path / "failed-run.xlsx"
    assert lo_runner.run(out, ["M_Main.RunFRTB_Silent"], input_book=source) == []
    result = openpyxl.load_workbook(out, data_only=True)
    errors = log_rows(result, "ERROR")
    assert any("Invalid numeric parameter RW (P_FX)" in str(row) for row in errors), errors
    assert not any("Finished with" in str(row) for row in log_rows(result, "INFO"))


def test_invariant_numeric_text_preserves_reconciled_outputs(lo_default_path, tmp_path):
    wb = openpyxl.load_workbook(lo_default_path)
    ws = wb["P_FX"]
    row = next(row for row in ws.iter_rows() if row[2].value == "RW")
    row[3].value = str(row[3].value)
    source = tmp_path / "numeric-text.xlsx"
    wb.save(source)
    out = tmp_path / "numeric-text-out.xlsx"
    assert lo_runner.run(out, ["M_Main.RunFRTB_Silent"], input_book=source) == []
    result = openpyxl.load_workbook(out, data_only=True)
    assert log_rows(result, "ERROR") == []
    assert reconcile(result) == []


def _set_table_value(wb, sheet, tag, key, column, value):
    """Set column (1 = first table column) of the row whose first cell is key in table [tag]."""
    in_table = False
    for row in wb[sheet].iter_rows(min_col=FIRST_COL):
        first = row[0].value
        if isinstance(first, str) and first.startswith("["):
            in_table = first.startswith(f"[{tag}]")
            continue
        if in_table and first is not None and str(first).strip().upper() == str(key).upper():
            row[column - 1].value = value
            return
    raise KeyError((sheet, tag, key))


INVALID_TABLE_VALUES = {
    "blank-risk-weight": ("P_CSR", "CSR_BUCKETS", 1, 6, None,
                          "[CSR_BUCKETS] data row 1, column 6 (P_CSR): blank"),
    "text-risk-weight": ("P_CSR", "CSR_BUCKETS", 1, 6, "0,5",
                         "[CSR_BUCKETS] data row 1, column 6 (P_CSR): '0,5'"),
    "blank-correlation": ("P_EQ", "EQ_GAMMA", 1, 3, None,
                          "[EQ_GAMMA] data row 1, column 3 (P_EQ): blank"),
    "text-liquidity-horizon": ("P_General", "VEGA_LH", "GIRR", 3, "60d",
                               "[VEGA_LH] data row 1, column 3 (P_General): '60d'"),
    "vertex-list": ("P_General", "VERTICES", "GIRR_TENORS", 2, "0.25,0.5,1,2,3,5,10,15,20,3O",
                    "Invalid number '3O' in GIRR_TENORS (P_General)"),
    "vertex-blank-entry": ("P_General", "VERTICES", "CSR_TENORS", 2, "0.5,1,,5,10",
                           "Blank entry in CSR_TENORS (P_General)"),
    "vertex-trailing-comma": ("P_General", "VERTICES", "COM_TENORS", 2, "0,0.25,0.5,1,2,3,5,10,15,20,30,",
                              "Blank entry in COM_TENORS (P_General)"),
    "vertex-empty": ("P_General", "VERTICES", "VEGA_MATURITIES", 2, None,
                     "Missing list VEGA_MATURITIES (P_General)"),
    # Fails in the vega step, after the delta step has written Out_Buckets rows
    "partial-outputs": ("P_General", "VEGA_LH", "GIRR", 1, "XX", "No liquidity horizon for GIRR"),
}


@pytest.mark.parametrize("case", sorted(INVALID_TABLE_VALUES))
def test_invalid_table_value_stops_the_run_and_clears_outputs(case, lo_default_path, tmp_path):
    sheet, tag, key, column, value, message = INVALID_TABLE_VALUES[case]
    wb = openpyxl.load_workbook(lo_default_path)
    _set_table_value(wb, sheet, tag, key, column, value)
    source = tmp_path / "bad-table.xlsx"
    wb.save(source)
    out = tmp_path / "failed-run.xlsx"
    assert lo_runner.run(out, ["M_Main.RunFRTB_Silent"], input_book=source) == []
    result = openpyxl.load_workbook(out, data_only=True)
    errors = log_rows(result, "ERROR")
    assert len(errors) == 1 and errors[0][1] == "RunFRTB" and message in errors[0][2], errors
    # The input book holds the default run's results; a failed run must not leave them behind
    for name in ("Out_Summary", "Out_RiskClass", "Out_Buckets", "Out_DRC"):
        assert sheet_rows(result, name) == [], name
    assert str(result["Results"].cell(5, FIRST_COL).value).startswith("Run failed (see the Log sheet): ")


def test_reference_numeric_text_matches_the_vba_contract():
    for text, value in {"0.5": 0.5, ".5": 0.5, "1.": 1.0, "-1.25e2": -125.0, "+2E-1": 0.2, " 0.5 ": 0.5}.items():
        assert ref._number(text) == value, text
    for text in ["abc", "0,5", "1,000", "1.2junk", "1e", "1e309", "1 2", "$1", "NaN", "inf", "1_000", "&H10", "1%", "", ".", "\t1"]:
        assert ref._number(text) is None, text
    assert ref._number(True) is None and ref._number(7) == 7.0


@pytest.mark.parametrize("key,clear_row", [(1, False), (7, False), (13, False), (7, True)])
def test_blank_parameter_key_or_internal_gap_fails(lo_default_path, tmp_path, key, clear_row):
    wb = openpyxl.load_workbook(lo_default_path)
    ws = wb["P_EQ"]
    tag_row = next(r for r in range(1, ws.max_row + 1)
                   if str(ws.cell(r, FIRST_COL).value).startswith("[EQ_GAMMA]"))
    row = tag_row + 1 + key
    for column in range(FIRST_COL, FIRST_COL + (14 if clear_row else 1)):
        ws.cell(row, column).value = None
    source = tmp_path / "blank-key.xlsx"
    wb.save(source)
    out = tmp_path / "blank-key-out.xlsx"
    assert lo_runner.run(out, ["M_Main.RunFRTB_Silent"], input_book=source) == []
    result = openpyxl.load_workbook(out, data_only=True)
    message = f"Blank {'data row' if clear_row else 'key'} in [EQ_GAMMA] data row {key} (P_EQ)"
    errors = log_rows(result, "ERROR")
    assert len(errors) == 1 and message in errors[0][2], errors
    for name in ("Out_Summary", "Out_RiskClass", "Out_Buckets", "Out_DRC"):
        assert sheet_rows(result, name) == [], name
    assert str(result["Results"].cell(5, FIRST_COL).value).startswith("Run failed")
