"""Reconciliation of the VBA engine (run in LibreOffice) with the reference.

The VBA modules are executed headless by tools/lo_runner.py; every output
cell is compared with tests/frtb_reference.py computed on the same
parameters and inputs. LibreOffice's VBA support is not Excel: these tests
show that the engine logic is right, not that Excel compiles or runs it.
"""

from __future__ import annotations

import math

import openpyxl
import pytest

import frtb_reference as ref
import lo_runner
import regulatory_defaults as reg

REL = 1e-9
ABS = 1e-6


def close(a, b):
    return math.isclose(float(a), float(b), rel_tol=REL, abs_tol=ABS)


FIRST_COL = 3  # tables start in column C
FIRST_DATA_ROW = 6  # header in row 5


def sheet_rows(wb, name, min_row=FIRST_DATA_ROW):
    rows = wb[name].iter_rows(min_row=min_row, min_col=FIRST_COL, values_only=True)
    return [r for r in rows if r and r[0] not in (None, "")]


def reconcile(wb):
    """Return a list of mismatch descriptions (empty when everything agrees)."""
    p = ref.params_from_workbook(wb)
    res = ref.run(p, ref.inputs_from_workbook(wb))
    problems = []

    rows = sheet_rows(wb, "Out_RiskClass")
    if len(rows) != len(ref.RISK_CLASSES) * len(ref.COMPONENTS) * len(ref.SCENARIOS):
        problems.append(f"Out_RiskClass has {len(rows)} rows")
    for rc, comp, scen, cap, _n, alt in (r[:6] for r in rows):
        if not close(cap, res["capital"][(rc, comp, scen)]):
            problems.append(f"capital {rc}/{comp}/{scen}: vba={cap} ref={res['capital'][(rc, comp, scen)]}")
        if (alt == "Y") != res["alt"][(rc, comp, scen)]:
            problems.append(f"alt flag {rc}/{comp}/{scen}")

    vba_buckets = {(r[0], r[1], r[2], str(r[3])): r[4:8] for r in sheet_rows(wb, "Out_Buckets")}
    if set(vba_buckets) != set(res["buckets"]):
        problems.append(f"bucket keys differ: {set(vba_buckets) ^ set(res['buckets'])}")
    for key, (kb, sb, n, note) in res["buckets"].items():
        v = vba_buckets.get(key)
        if v is None:
            continue
        if not (close(v[0], kb) and close(v[1], sb) and v[2] == n and (v[3] or "") == note):
            problems.append(f"bucket {key}: vba={v} ref={(kb, sb, n, note)}")

    summary = {r[0]: r[1] for r in sheet_rows(wb, "Out_Summary")}
    for key, value in res["summary"].items():
        if isinstance(value, str):
            if summary.get(key) != value:
                problems.append(f"summary {key}: vba={summary.get(key)} ref={value}")
        elif not close(summary.get(key), value):
            problems.append(f"summary {key}: vba={summary.get(key)} ref={value}")

    drc_rows = sheet_rows(wb, "Out_DRC")
    vba_obl = {r[0]: r[1:6] for r in drc_rows}
    if set(vba_obl) != set(res["drc_obligors"]):
        problems.append("DRC obligor sets differ")
    for name, (bucket, quality, rw, lg, sh) in res["drc_obligors"].items():
        v = vba_obl.get(name)
        if v and not (v[0] == bucket and v[1] == quality and close(v[2], rw) and close(v[3], lg) and close(v[4], sh)):
            problems.append(f"DRC obligor {name}: vba={v} ref={(bucket, quality, rw, lg, sh)}")
    vba_drc_buckets = {
        r[7]: r[8:14]
        for r in wb["Out_DRC"].iter_rows(min_row=FIRST_DATA_ROW, min_col=FIRST_COL, values_only=True)
        if len(r) > 7 and r[7]
    }
    for b, vals in res["drc_buckets"].items():
        v = vba_drc_buckets.get(b)
        if v is None or not all(close(x, y) for x, y in zip(v, vals)):
            problems.append(f"DRC bucket {b}: vba={v} ref={vals}")
    return problems


def log_rows(wb, level=None):
    rows = sheet_rows(wb, "Log")
    return [r for r in rows if level is None or r[0] == level]


# --------------------------------------------------------------------------
# Default sample portfolio
# --------------------------------------------------------------------------
def test_default_run_reconciles_with_reference(lo_default_run):
    assert reconcile(lo_default_run) == []


def test_default_run_has_clean_log(lo_default_run):
    assert log_rows(lo_default_run, "ERROR") == []
    assert log_rows(lo_default_run, "WARN") == []


def test_default_run_exercises_alternative_sb(lo_default_run):
    alt = [r for r in sheet_rows(lo_default_run, "Out_RiskClass") if r[5] == "Y"]
    assert alt, "the sample portfolio should trigger the alternative Sb specification"


def test_in_workbook_regression_test_passes(lo_default_run):
    ws = lo_default_run["Test_Expected"]
    assert str(ws["C3"].value).startswith("Result: PASS")
    statuses = [r[4] for r in ws.iter_rows(min_row=7, min_col=FIRST_COL, values_only=True) if r[0]]
    assert len(statuses) == 52 and set(statuses) == {"PASS"}


def test_workbook_layout(lo_default_run):
    expected = [
        "README", "Control", "Results", "In_Delta", "In_Vega", "In_Curvature", "In_DRC", "In_RRAO",
        "Out_Summary", "Out_RiskClass", "Out_Buckets", "Out_DRC", "Log", "P_General", "P_GIRR", "P_CSR",
        "P_EQ", "P_COM", "P_FX", "P_DRC", "P_RRAO", "Test_Expected",
    ]
    assert lo_default_run.sheetnames == expected


# --------------------------------------------------------------------------
# Parameters written by M_Setup vs the independently typed defaults
# --------------------------------------------------------------------------
def test_parameter_sheets_match_regulatory_defaults(lo_default_run):
    got = ref.params_from_workbook(lo_default_run)
    exp = reg.build_params()
    assert got.girr_rw == pytest.approx(exp.girr_rw)
    assert got.girr == pytest.approx(exp.girr)
    assert got.girr_specified == exp.girr_specified
    assert got.csr_buckets == exp.csr_buckets
    assert got.csr == pytest.approx(exp.csr)
    assert got.csr_sector == pytest.approx(exp.csr_sector)
    assert set(got.eq_buckets) == set(exp.eq_buckets)
    for b in exp.eq_buckets:
        assert got.eq_buckets[b]["type"] == exp.eq_buckets[b]["type"]
        assert got.eq_buckets[b]["rw_spot"] == pytest.approx(exp.eq_buckets[b]["rw_spot"])
        assert got.eq_buckets[b]["rw_repo"] == pytest.approx(exp.eq_buckets[b]["rw_repo"])
        assert got.eq_buckets[b]["rho"] == pytest.approx(exp.eq_buckets[b]["rho"])
    assert got.eq_gamma == pytest.approx(exp.eq_gamma)
    assert set(got.com_buckets) == set(exp.com_buckets)
    for b, v in exp.com_buckets.items():
        assert got.com_buckets[b]["rw"] == pytest.approx(v["rw"])
        assert got.com_buckets[b]["rho"] == pytest.approx(v["rho"])
    assert got.com_gamma == pytest.approx(exp.com_gamma)
    assert got.fx == pytest.approx(exp.fx)
    assert got.fx_pairs == exp.fx_pairs
    assert got.drc_rw == pytest.approx(exp.drc_rw)
    assert got.drc_seniority == exp.drc_seniority
    assert got.drc == pytest.approx(exp.drc)
    assert got.drc_buckets == exp.drc_buckets
    assert got.rrao_rw == pytest.approx(exp.rrao_rw)
    assert got.vertices == exp.vertices
    assert sorted(got.vega_lh) == sorted(exp.vega_lh)
    assert (got.high_mult, got.low_mult, got.vega_rw_sigma, got.vega_alpha) == (1.25, 0.75, 0.55, 0.01)


# --------------------------------------------------------------------------
# Discretions switched on and invalid input rows
# --------------------------------------------------------------------------
INVALID_ROWS = {
    "In_Delta": [
        ["X1", "CSR_SEC_CTP", 1, "A", "BOND", 5, 1000],
        ["X2", "GIRR", "EUR", "ESTR", "YIELD", 7, 1000],
        ["X3", "FX", "EUR", "EUR", None, None, 1000],
        ["X4", "EQ", 99, "A", "SPOT", None, 1000],
        ["X5", "EQ", 8, "A", "SPOT", None, "abc"],
        ["X6", "COM", 2, "BRENT", "ROTTERDAM", "abc", 1000],
        ["X7", "EQ", 8.6, "A", "SPOT", None, 1000],
    ],
    "In_Vega": [["Y1", "EQ", 8, "A", None, 2, None, 1000]],
    "In_Curvature": [["Z1", "EQ", 8, "A", None, 1000, None]],
    "In_DRC": [
        ["K1", "A", "CORP", "ZZZ", "SENIOR", 1, 1, 1, None],
        ["K2", "A", "XXX", "A", "SENIOR", 1, 1, 1, None],
        ["K3", "A", "CORP", "A", "JUNK", 1, 1, 1, None],
        ["K4", "BANK_A", "CORP", "BBB", "SENIOR", 100, 100, 1, None],
        ["K5", "BANK_C", "CORP", "A", "SENIOR", 1000, 1000, "abc", None],
        ["K6", "BANK_C", "CORP", "A", "SENIOR", 1000, 1000, 1, "x"],
    ],
    "In_RRAO": [["Q1", "FOO", 1000], ["Q2", "OTHER", -5]],
}


@pytest.fixture(scope="module")
def lo_variant_run(lo_default_path, tmp_path_factory):
    wb = openpyxl.load_workbook(lo_default_path)
    ws = wb["Control"]
    for row in ws.iter_rows(min_row=1, min_col=FIRST_COL):
        if row[0].value in ("GIRR_SPECIFIED_CCY_DISCRETION", "FX_SPECIFIED_PAIR_DISCRETION", "FX_CURVATURE_DIVISOR"):
            row[1].value = True
    for sheet, rows in INVALID_ROWS.items():
        ws = wb[sheet]
        r = FIRST_DATA_ROW
        while ws.cell(r, FIRST_COL).value not in (None, "") or ws.cell(r, FIRST_COL + 1).value not in (None, ""):
            r += 1
        for values in rows:
            for c, v in enumerate(values, start=FIRST_COL):
                ws.cell(r, c).value = v
            r += 1
    folder = tmp_path_factory.mktemp("lo-variant")
    source = folder / "variant-in.xlsx"
    wb.save(source)
    out = folder / "variant-out.xlsx"
    errors = lo_runner.run(out, ["M_Main.RunFRTB_Silent"], input_book=source)
    assert errors == [], errors
    return openpyxl.load_workbook(out, data_only=True)


def test_variant_run_reconciles_with_reference(lo_variant_run):
    assert reconcile(lo_variant_run) == []


def test_variant_run_logs_every_invalid_row(lo_variant_run):
    errors = log_rows(lo_variant_run, "ERROR")
    sources = sorted(r[1] for r in errors)
    assert len(errors) == 16, errors
    assert sum(s.startswith("In_Delta") for s in sources) == 7
    assert sum(s.startswith("In_DRC") for s in sources) == 5
    warnings = log_rows(lo_variant_run, "WARN")
    assert len(warnings) == 1 and "credit qualities" in warnings[0][2]


def test_variant_run_applies_discretions(lo_variant_run, lo_default_run):
    def cap(wb, rc, comp):
        return {(r[0], r[1], r[2]): r[3] for r in sheet_rows(wb, "Out_RiskClass")}[(rc, comp, "MEDIUM")]

    # Specified currency and pair discretions lower the GIRR and FX delta charges
    # (all sample FX currencies are specified against EUR, directly or as crosses)
    assert cap(lo_variant_run, "GIRR", "DELTA") < cap(lo_default_run, "GIRR", "DELTA")
    assert cap(lo_variant_run, "FX", "DELTA") == pytest.approx(cap(lo_default_run, "FX", "DELTA") / math.sqrt(2))


# --------------------------------------------------------------------------
# Empty inputs
# --------------------------------------------------------------------------
def test_empty_inputs_give_zero_capital(lo_default_path, tmp_path):
    wb = openpyxl.load_workbook(lo_default_path)
    for sheet in ("In_Delta", "In_Vega", "In_Curvature", "In_DRC", "In_RRAO"):
        ws = wb[sheet]
        ws.delete_rows(FIRST_DATA_ROW, ws.max_row)
    source = tmp_path / "empty-in.xlsx"
    wb.save(source)
    out = tmp_path / "empty-out.xlsx"
    assert lo_runner.run(out, ["M_Main.RunFRTB_Silent"], input_book=source) == []
    result = openpyxl.load_workbook(out, data_only=True)
    summary = {r[0]: r[1] for r in sheet_rows(result, "Out_Summary")}
    assert summary["TOTAL"] == 0 and summary["ERRORS"] == 0
    assert all(r[3] == 0 for r in sheet_rows(result, "Out_RiskClass"))
    assert sheet_rows(result, "Out_Buckets") == []
